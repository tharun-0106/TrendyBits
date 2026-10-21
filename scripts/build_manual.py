#!/usr/bin/env python
"""Build docs/TrendyBits_App_Manual_and_Guide.docx (python-docx) from the code base and the measured telemetry.

Run AFTER the evaluation scripts so the tables contain real numbers:
  python scripts/run_all.py && python scripts/build_manual.py
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

DATA = ROOT / "data"
OUT_DIR = ROOT / "docs"
FIG = OUT_DIR / "figures"
OUT = OUT_DIR / "TrendyBits_App_Manual_and_Guide.docx"
ACCENT = RGBColor(0xC8, 0x10, 0x2E)


def load(name: str) -> dict:
    p = DATA / name
    return json.loads(p.read_text()) if p.exists() else {}


CLF, PIPE, ABL, LLM, XAIV = (load(n) for n in ("model_metrics.json", "pipeline_model_metrics.json", "ablation_metrics.json",
                                                 "llm_test_metrics.json", "xai_verification.json"))
MAN = {n: load(f"datasets/{n}/manifest.json") for n in ("newssumm", "liar", "integrated")}


def f(v, d=3):
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        return f"{v:.{d}f}"
    return str(v)


# =============================================================================== figures
def fig_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 6.2))
    ax.set_xlim(0, 110)
    ax.set_ylim(0, 62)
    ax.axis("off")

    def box(x, y, w, h, text, fc, ec="#333", fs=8.5, bold=False):
        ax.add_patch(plt.Rectangle((x, y), w, h, fc=fc, ec=ec, lw=1.2, joinstyle="round"))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, wrap=True,
                fontweight="bold" if bold else "normal")

    def arrow(x1, y1, x2, y2, c="#444"):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", color=c, lw=1.3))

    box(2, 50, 26, 9, "React 18 + Vite UI\nFeed · Article · Fact-Check\nChatbot · Admin & XAI", "#fde8eb", bold=True)
    box(34, 50, 30, 9, "FastAPI (backend/app/main.py)\n/api/chat /search /summarize\n/api/analyze /admin/metrics /xai", "#e0f2fe", bold=True)
    box(70, 50, 36, 9, "SQLite trendybits.db\narticles · sources · users · queries\nreports · eval logs · feedback · security", "#f3f4f6")
    arrow(28, 54.5, 34, 54.5)
    arrow(64, 54.5, 70, 54.5)
    # Pipeline A
    ax.add_patch(plt.Rectangle((1, 2), 50, 44, fc="none", ec="#0ea5e9", lw=2, ls="--"))
    ax.text(26, 43.5, "PIPELINE A — Personalised retrieval & faithful summarisation", ha="center", fontsize=9, color="#0369a1", fontweight="bold")
    box(3, 33, 22, 8, "Module 1\nIntent extraction\n(qwen:4b + fallback)", "#e0f2fe")
    box(27, 33, 22, 8, "Module 2\nScraper (RSS/URL)\n+ dedup → SQLite", "#e0f2fe")
    box(3, 21, 46, 9, "Module 2  Sentence-BERT all-MiniLM-L6-v2 → FAISS IndexFlatIP\n+ BM25 hybrid · topic/source/time filters · profile boost · MMR re-rank", "#e0f2fe")
    box(3, 6, 46, 11, "Module 3  Faithful summariser (qwen:4b)\nchunked map-reduce · entity/number/hedge guards\nNO true/false verdict · source URL always attached", "#e0f2fe")
    arrow(14, 33, 14, 30)
    arrow(26, 21, 26, 17)
    # Pipeline B
    ax.add_patch(plt.Rectangle((55, 2), 54, 44, fc="none", ec=ACCENT_HEX, lw=2, ls="--"))
    ax.text(82, 43.5, "PIPELINE B — Independent credibility analysis", ha="center", fontsize=9, color=ACCENT_HEX, fontweight="bold")
    box(57, 35, 24, 6.5, "Module 5 Security guard\n(injection strip / toxicity)", "#fee2e2")
    box(83, 35, 24, 6.5, "Step 1 Claim decomposition\n(fact / opinion / CTA / media)", "#fee2e2")
    box(57, 26, 24, 7, "Step 2 RAG evidence\n(trusted news + fact-checks)", "#fee2e2")
    box(83, 26, 24, 7, "Step 3 BERT-LIAR 6-class\n+ ISOT-Bharat binary signals", "#fee2e2")
    box(57, 16, 50, 8, "Step 4 Evidence reasoner: deepseek-r1:8b / gemma2:9b (NLI fallback)\nSUPPORTED · FALSE · MISLEADING · UNVERIFIED\nconflict, credibility & recency weighting", "#fee2e2", fs=7.8)
    box(57, 4, 24, 10, "Step 5 Report\nsource trust · fact accuracy\noverall estimate · timeline", "#fee2e2")
    box(83, 4, 24, 10, "Module 6 XAI + telemetry\nLIME token heatmap\nexact SHAP over 5 dims", "#ede9fe")
    arrow(81, 38, 83, 38)
    arrow(95, 35, 95, 33)
    arrow(69, 35, 69, 33)
    arrow(69, 26, 69, 24)
    arrow(95, 26, 95, 24)
    arrow(69, 16, 69, 14)
    arrow(95, 16, 95, 14)
    ax.text(55, 47.6, "the two pipelines share only article IDs and source URLs", ha="center", fontsize=8, color="#555", style="italic")
    fig.tight_layout()
    fig.savefig(path, dpi=170)
    plt.close(fig)


ACCENT_HEX = "#c8102e"


def fig_cm(matrix, labels, title, path: Path) -> None:
    import numpy as np

    m = np.array(matrix)
    fig, ax = plt.subplots(figsize=(4.6 if len(labels) > 2 else 3.2, 3.9 if len(labels) > 2 else 2.9))
    ax.imshow(m, cmap="Reds")
    ax.set_xticks(range(len(labels)), labels, rotation=40, ha="right", fontsize=7)
    ax.set_yticks(range(len(labels)), labels, fontsize=7)
    for i in range(m.shape[0]):
        for j in range(m.shape[1]):
            ax.text(j, i, m[i, j], ha="center", va="center", fontsize=8, color="white" if m[i, j] > m.max() * 0.55 else "black")
    ax.set_xlabel("predicted", fontsize=8)
    ax.set_ylabel("gold", fontsize=8)
    ax.set_title(title, fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=170)
    plt.close(fig)


def fig_ablation(path: Path) -> None:
    rows = ABL.get("rows", [])
    if not rows:
        return
    metrics = [("precision_at_5", "P@5"), ("ild", "ILD"), ("unique_source_ratio", "Uniq. sources"), ("verdict_accuracy", "Verdict acc."),
               ("credibility_band_accuracy", "Cred. band acc."), ("rougeL", "ROUGE-L")]
    import numpy as np

    fig, ax = plt.subplots(figsize=(10, 3.8))
    x = np.arange(len(metrics))
    w = 0.12
    colors = ["#111827", "#0ea5e9", "#22c55e", "#a855f7", "#f97316", "#ef4444", "#eab308"]
    for i, r in enumerate(rows):
        ax.bar(x + (i - 3) * w, [r[k] or 0 for k, _ in metrics], w, label=r["configuration"].split(":")[0], color=colors[i % 7])
    ax.set_xticks(x, [m for _, m in metrics], fontsize=8)
    ax.set_ylim(0, 1.05)
    ax.legend(fontsize=7, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.12))
    ax.set_title("Ablation: complete system vs Models 1–6 (identical held-out samples)", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=170)
    plt.close(fig)


def fig_xai_example(path_shap: Path, path_lime: Path) -> dict:
    from app.services import get_services

    rep = get_services().analyze_article("tb-001", save=False)
    shap = rep["xai_explanation"]["shap"]
    fig, ax = plt.subplots(figsize=(7, 2.8))
    labels = [c["label"] for c in shap["contributions"]]
    vals = [c["contribution"] for c in shap["contributions"]]
    ax.barh(labels, vals, color=["#16a34a" if v >= 0 else "#dc2626" for v in vals])
    ax.axvline(0, color="#333", lw=0.8)
    lo, hi = min(0, min(vals)), max(0, max(vals))
    pad = (hi - lo) * 0.18 + 1
    ax.set_xlim(lo - pad, hi + pad)
    ax.invert_yaxis()
    for i, v in enumerate(vals):
        ax.text(v + (0.4 if v >= 0 else -0.4), i, f"{v:+.2f}", va="center", ha="left" if v >= 0 else "right", fontsize=8)
    ax.set_title(f"SHAP contributions — baseline {shap['baseline_score']} → final {shap['final_score']} (tb-001)", fontsize=9)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(path_shap, dpi=170)
    plt.close(fig)
    lime = rep["xai_explanation"]["lime"]
    toks = lime["top_negative"][:6][::-1] + lime["top_positive"][:4]
    fig, ax = plt.subplots(figsize=(7, 3))
    ax.barh([t["token"] for t in toks], [t["weight"] for t in toks], color=["#16a34a" if t["weight"] > 0 else "#dc2626" for t in toks])
    ax.axvline(0, color="#333", lw=0.8)
    ax.set_title(f"LIME token weights (R² = {lime['local_fidelity_r2']}, {lime['num_samples']} perturbations)", fontsize=9)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(path_lime, dpi=170)
    plt.close(fig)
    return rep


# =============================================================================== docx helpers
def set_cell_bg(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(doc, header: list[str], rows: list[list], widths: list[float] | None = None, font: float = 8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(c, "C8102E")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(f(v) if not isinstance(v, str) else v)
            r.font.size = Pt(font)
            if v in ("PASS",):
                set_cell_bg(cells[i], "DCFCE7")
            elif v in ("FAIL",):
                set_cell_bg(cells[i], "FEE2E2")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def code(doc, text: str) -> None:
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(8.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F3F4F6")
        pPr.append(shd)
    doc.add_paragraph()


def eq(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Cambria Math"
    r.italic = True
    r.font.size = Pt(11)


def bullets(doc, items: list[str]) -> None:
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if "::" in it:
            head, tail = it.split("::", 1)
            p.add_run(head.strip()).bold = True
            p.add_run(" — " + tail.strip())
        else:
            p.add_run(it)


def para(doc, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)


def picture(doc, path: Path, width_cm: float, caption: str) -> None:
    if path.exists():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(8.5)


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        elif text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Right-click → Update Field (or press F9) to build the table of contents."
            run._r.append(t)


# =============================================================================== document
def build() -> Path:
    OUT_DIR.mkdir(exist_ok=True)
    FIG.mkdir(exist_ok=True)
    fig_architecture(FIG / "architecture.png")
    if CLF.get("liar"):
        m = CLF["liar"]["primary"]
        fig_cm(m["confusion_matrix"], m["labels"], "BERT-LIAR 6-class confusion matrix", FIG / "cm_liar.png")
        m = CLF["integrated"]["primary"]
        fig_cm(m["confusion_matrix"], ["False (0)", "True (1)"], "ISOT+Bharat binary", FIG / "cm_integrated.png")
    fig_ablation(FIG / "ablation.png")
    fig_xai_example(FIG / "shap_example.png", FIG / "lime_example.png")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for lvl, size in ((1, 17), (2, 13), (3, 11.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.color.rgb = ACCENT if lvl == 1 else RGBColor(0x1B, 0x1B, 0x1F)
        h.font.size = Pt(size)
    footer = sec.footer.paragraphs[0]
    footer.text = "TrendyBits — App Manual & Guide · Pipeline A never judges truth · Pipeline B cites only retrieved evidence"
    footer.runs[0].font.size = Pt(8)

    # ---------------- title page
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("TB  TrendyBits")
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = ACCENT
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run("LLM-Powered Conversational Platform for Personalized News\nand Fake-News & Multi-Dimensional Credibility Analysis").font.size = Pt(14)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run("Application Manual, Architecture, Explainable-AI Algorithm, Model Performance & Operations Guide").italic = True
    doc.add_paragraph()
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run(f"Version 1.0 · generated {date.today().isoformat()} from the repository's measured telemetry").font.size = Pt(10)
    page_break(doc)

    doc.add_heading("Contents", 1)
    add_toc(doc)
    page_break(doc)

    # ---------------- 1. overview
    doc.add_heading("1. Overview and reading guide", 1)
    para(doc, "TrendyBits is a text-first NLP application motivated by the FND-LLM framework: a specialised supervised classifier "
              "(SLM/BERT) supplies calibrated signals, while LLMs provide intent understanding, faithful summarisation and "
              "evidence-grounded reasoning. The platform is split into two connected but strictly independent pipelines:")
    bullets(doc, [
        "Pipeline A — Personalised News Retrieval & Faithful Summarisation:: conversational intent extraction, live scraping, "
        "Sentence-BERT + FAISS semantic search with personalisation and MMR diversity, and a qwen:4b summariser that is forbidden "
        "from judging truth.",
        "Pipeline B — Independent Fake-News & Multi-Dimensional Credibility Analysis:: claim decomposition, retrieval-augmented "
        "evidence search, BERT classifier signals, an 8B–9B evidence reasoner (deepseek-r1:8b / gemma2:9b) and a transparent, "
        "multi-dimensional report with LIME + SHAP explanations and live model telemetry.",
    ])
    para(doc, "The two pipelines share only article IDs and source URLs. Pipeline A's code never imports Pipeline B (this is unit-tested), "
              "so fluent generated text can never be presented as evidence of truth, and binary Real/Fake labels are replaced by a "
              "graded report whose every citation must come from the retrieved evidence set.")
    doc.add_heading("1.1 Important: what the shipped numbers mean", 2)
    para(doc, "Every metric in this manual was measured by running the repository's own scripts. To make the repository runnable "
              "immediately after cloning, the shipped run used:", None)
    bullets(doc, [
        "Synthetic, schema-faithful benchmark partitions:: template-generated NewsSumm / LIAR / Integrated datasets (with label noise) "
        "and a 26-article demo seed corpus. They are NOT the IEEE Dataport data; import the real files (Section 11.4) and re-run to "
        "obtain publishable numbers.",
        "The deterministic fallback engine:: Ollama was not running, so summarisation, intent extraction and evidence reasoning used the "
        "built-in deterministic NLP engine; embeddings used the TF-IDF + LSA fallback (sentence-transformers / FAISS not installed); "
        "the classifiers used the fast TF-IDF + Logistic-Regression back-end (bert-base-uncased fine-tuning is available with --mode bert).",
        "Development-set caveat:: the 24-claim evidence benchmark and the 18 test cases were authored together with the deterministic "
        "engine, so their near-perfect scores are pipeline sanity checks, not estimates of real-world accuracy.",
    ])
    eng = LLM.get("engine", {})
    table(doc, ["Component", "Back-end used for the shipped numbers"], [
        ["Retriever (Model 1)", eng.get("retriever", "—")],
        ["Summariser / intent (Model 2)", eng.get("summarizer", "—")],
        ["LIAR & Integrated classifiers (Models 3, 4)", eng.get("classifiers", "—")],
        ["Evidence reasoner (Model 5)", ", ".join(eng.get("reasoner", [])) or "—"],
        ["LLM-as-a-judge", eng.get("judge", "—")],
        ["Ollama reachable during the run", str(eng.get("ollama_reachable"))],
    ], [6, 11])

    # ---------------- 2. architecture
    page_break(doc)
    doc.add_heading("2. System architecture", 1)
    picture(doc, FIG / "architecture.png", 17, "Figure 1 — TrendyBits architecture: Pipeline A (blue) and Pipeline B (red) share only article IDs and URLs.")
    doc.add_heading("2.1 Technology stack", 2)
    table(doc, ["Layer", "Technology", "Fallback when unavailable"], [
        ["Backend", "Python 3.11, FastAPI, Uvicorn, Pydantic v2", "—"],
        ["Database", "SQLite (data/trendybits.db): articles, sources, users, queries, reports, model_eval_logs, feedback, security_events", "—"],
        ["Embeddings / index", "sentence-transformers all-MiniLM-L6-v2 + faiss-cpu IndexFlatIP", "TF-IDF(1–2-gram) ⊕ LSA (TruncatedSVD) + concept lexicon; NumPy cosine"],
        ["Pipeline A LLM", "Ollama qwen:4b (or qwen2.5:3b), ≈2.5 GB VRAM", "Deterministic intent engine + extractive salience summariser"],
        ["Pipeline B LLM", "Ollama deepseek-r1:8b or gemma2:9b (≈5 GB VRAM); llama3.1:8b selectable", "Deterministic NLI reasoner (clause-level stance rules)"],
        ["Judge LLM", "Ollama mistral:7b via CustomOllamaJudge (DeepEval)", "Deterministic faithfulness / relevancy scorers"],
        ["Classifiers", "PyTorch + HuggingFace bert-base-uncased (AutoModelForSequenceClassification)", "TF-IDF (+credit history) → Logistic Regression; TF-IDF→SVD→MLP baseline"],
        ["XAI", "Custom LIME perturbation engine + exact Shapley values (backend/app/xai_explainer.py)", "— (pure NumPy / scikit-learn)"],
        ["Scraping", "feedparser, requests, beautifulsoup4, newspaper3k", "stdlib XML RSS parser; BeautifulSoup heuristics"],
        ["Testing", "pytest, deepeval, ragas, rouge-score, scikit-learn", "built-in ROUGE, deterministic scorers"],
        ["Frontend", "React 18 + Vite 5, hand-written CSS (light/dark tokens)", "—"],
        ["Deployment", "Dockerfiles for backend & frontend (nginx), docker-compose with optional Ollama profile", "—"],
    ], [3, 7.5, 6.5])
    doc.add_heading("2.2 Repository layout", 2)
    code(doc, """
TrendyBits/
├── backend/
│   ├── app/
│   │   ├── main.py            FastAPI routes (Pipeline A, Pipeline B, XAI, admin)
│   │   ├── services.py        service container wiring DB → retriever → analyzer
│   │   ├── config.py          env-driven configuration (models, paths, knobs)
│   │   ├── database.py        SQLite schema + data access
│   │   ├── seed_data.py       source registry + 26 synthetic demo articles
│   │   ├── datasets.py        IEEE Dataport schemas, loaders, generators, audit, splits
│   │   ├── text_utils.py      tokenisation, NER, quantities, hedges, taxonomy
│   │   ├── embeddings.py      SBERT / TF-IDF+LSA embedder, FAISS / NumPy index
│   │   ├── retrieval.py       hybrid search, filters, personalisation, MMR, IR metrics
│   │   ├── scraper.py         RSS + URL scraping, cleaning, dedup, SSRF guard
│   │   ├── llm_service.py     Ollama client, Module 1 intent, Module 3 summariser
│   │   ├── security.py        Module 5 injection guard + toxicity neutraliser
│   │   ├── classifiers.py     Models 3/4 training, metrics and inference
│   │   ├── credibility.py     Module 4 Pipeline B analyser and report
│   │   ├── xai_explainer.py   Module 6 LIME + SHAP
│   │   ├── metrics_store.py   Module 6 telemetry merge for /api/admin/metrics
│   │   ├── judge.py           CustomOllamaJudge, ROUGE, faithfulness, relevancy
│   │   └── evaluation.py      ablation (Table 1) + 18 test cases (Tables 2–5)
│   ├── requirements.txt, requirements-ml.txt, Dockerfile
├── frontend/  (React + Vite: src/App.jsx, styles.css, components/, pages/, Dockerfile, nginx.conf)
├── scripts/   import_ieee_dataport.py, train_bert_classifier.py, run_ablation_study.py,
│              run_llm_test_suite.py, verify_xai.py, check_ollama.py, run_all.py, build_manual.py
├── tests/     test_pipelines.py, test_api.py, test_deepeval_suite.py, conftest.py
├── data/      datasets/, model_metrics.json, ablation_metrics.json, llm_test_metrics.json, …
├── models/    liar_fast.joblib, integrated_fast.joblib (bert_*/ after --mode bert)
├── docs/      this manual + figures
└── docker-compose.yml, README.md, .env.example
""")

    # ---------------- 3. datasets
    page_break(doc)
    doc.add_heading("3. Datasets and the IEEE Dataport ingestion pipeline", 1)
    para(doc, "Sources: TrendyBits News Summarization Dataset (https://ieee-dataport.org/documents/trendybits-news-summarization-dataset) and "
              "TrendyBits Fake News Dataset (https://ieee-dataport.org/documents/trendybits-fake-news-dataset). IEEE Dataport requires a "
              "signed-in download, so files are downloaded manually and passed to scripts/import_ieee_dataport.py. Without files the script "
              "generates schema-identical synthetic partitions.")
    table(doc, ["Dataset", "Unit", "Exact columns / labels"], [
        ["1. NewsSumm", "news article", "Source, Content, Summary (publishers: The Hindu, Indian Express, Hindustan Times, Times of India, CNN, Daily Mail, Reuters, Multi-News, XSum)"],
        ["2. LIAR (PolitiFact)", "short claim", "id, label, statement, subject, speaker, job_title, state, party, barely_true_counts, false_counts, half_true_counts, mostly_true_counts, pants_on_fire_counts, context · labels: pants-fire, false, barely-true, half-true, mostly-true, true"],
        ["3. Integrated ISOT + BharatFakeNewsKosh (EN)", "full article", "date, title, subject, text, label (1 = True, 0 = False)"],
    ], [3.5, 2.5, 11])
    para(doc, "Because LIAR (6-class, claim level) and the Integrated set (binary, article level) differ in unit of analysis and label design, "
              "they are trained, evaluated and reported separately (Models 3 and 4).")
    doc.add_heading("3.1 Audit and leakage-free splitting", 2)
    bullets(doc, ["Missing required fields are dropped; labels are normalised (e.g. 'Pants on Fire' → pants-fire, 'fake' → 0).",
                  "Non-English rows are removed (ASCII-letter ratio ≥ 0.85; Devanagari rows dropped).",
                  "Exact duplicates (normalised text) are removed BEFORE splitting, preventing train/test leakage.",
                  "Stratified 80/10/10 train/valid/test splits (seed 42) are created before any fitting or prompt tuning. Hyper-parameters "
                  "(e.g. logistic-regression C) are tuned on VALID only; the TEST partition is fingerprinted with SHA-256 of its row ids "
                  "in manifest.json and every metric file records that fingerprint."])
    rows = []
    for n, m in MAN.items():
        a = m.get("audit", {})
        rows.append([n, m.get("origin", "—")[:40], a.get("rows_in"), a.get("removed_missing"), a.get("removed_non_english"),
                     a.get("removed_duplicates"), f"{m['sizes']['train']}/{m['sizes']['valid']}/{m['sizes']['test']}" if m.get("sizes") else "—",
                     (m.get("test_sha256") or "")[:12] + "…"])
    table(doc, ["Dataset", "Origin", "Rows in", "Missing", "Non-EN", "Dupes", "Train/Valid/Test", "Test SHA-256"], rows,
          [2.2, 3.6, 1.4, 1.4, 1.4, 1.3, 2.6, 2.6], font=7.5)

    # ---------------- 4. modules
    page_break(doc)
    doc.add_heading("4. Module-by-module walkthrough", 1)
    doc.add_heading("4.1 Module 1 — Conversational intent extraction (llm_service.extract_intent)", 2)
    para(doc, "The user query is first sanitised by the security guard, then the deterministic engine extracts intent and slots; when qwen:4b "
              "is reachable its JSON output (schema-validated) overrides individual slots. Intents: news_retrieval, clarify_or_top_stories, "
              "fact_check, out_of_scope. Slots: topic (8-topic taxonomy), entities (capitalised-phrase NER + gazetteer, minus publisher "
              "names), source_filter (publisher alias table), time_window (today, yesterday, this week, last week, this month, recent).")
    bullets(doc, ["Ambiguous queries (TC-02) such as \"what's happening?\" with a sparse profile return a clarifying question plus "
                  "diverse top stories (MMR with λ = 0.3).",
                  "Out-of-scope requests (TC-03) — medical dosage, investment, legal, coding — are declined politely with no articles.",
                  "fact_check intents are routed to Pipeline B; the chatbot only links to the independent report."])
    doc.add_heading("4.2 Module 2 — Scraping, semantic search and MMR (scraper.py, retrieval.py)", 2)
    para(doc, "The scraper reads RSS feeds (The Hindu, Indian Express, Times of India, BBC, Reuters) or any public URL (SSRF-guarded), extracts "
              "title/body/date/author, assigns a topic, scans for embedded instructions and de-duplicates on URL and normalised content hash. "
              "The retriever scores every document as:")
    eq(doc, "rel(Dᵢ) = cos(e_Q, e_Dᵢ) + 0.15·BM25ₙ(Q, Dᵢ) + 0.12·[topic ∈ P] + 0.08·[source ∈ P] + 0.05·Σ[entity ∈ Dᵢ] + 0.05·e^(−age/14)")
    para(doc, "Topic, source and time-window filters are applied as hard filters with graceful degradation to soft boosts when nothing matches. "
              "The top 30 candidates are re-ranked with Maximal Marginal Relevance plus a same-publisher penalty γ = 0.15:")
    eq(doc, "MMR = argmax_{Dᵢ ∈ R∖S} [ λ·Sim₁(Dᵢ, Q) − (1−λ)·max_{Dⱼ ∈ S} Sim₂(Dᵢ, Dⱼ) − γ·|{Dⱼ ∈ S : pub(Dⱼ) = pub(Dᵢ)}| ],  λ = 0.7")
    para(doc, "Intra-List Diversity is ILD = mean over pairs (1 − cos(Dᵢ, Dⱼ)); the unique-source ratio is |publishers| / K.")
    doc.add_heading("4.3 Module 3 — Faithful summariser, Pipeline A (llm_service.summarize)", 2)
    bullets(doc, ["Untrusted text is sanitised (embedded instructions removed) and toxic / graphic quotations are replaced by a single neutral placeholder (TC-16).",
                  "Chunked map-reduce: the body is split into sentence-aligned chunks (350 words in fallback mode, 1,400 words for qwen:4b); every chunk "
                  "contributes key points, so the tail of long articles is never silently truncated (TC-07).",
                  "Guards applied to every summary (LLM or fallback): no invented numbers, hedges kept whenever a hedged sentence is paraphrased "
                  "(TC-05), no certainty upgrades, and no system-authored true/false verdict (TC-06). If an LLM summary fails any guard it is "
                  "rejected and the deterministic summary is used.",
                  "Every summary carries the original headline, publisher, publication date and URL (TC-18)."])
    doc.add_heading("4.4 Module 4 — Independent credibility analyser, Pipeline B (credibility.py)", 2)
    table(doc, ["Step", "What happens"], [
        ["1 Pre-processing & claim decomposition", "Sanitise title/body/user note; split sentences; label each as factual claim, opinion, call-to-action, question, satire or background; rank by check-worthiness (quantities, entities, absolute quantifiers, assertive verbs); add image captions as media claims."],
        ["2 Evidence retrieval (RAG)", "Per claim, hybrid search restricted to news / official / fact-check / blog items (viral items under review are never evidence); best two sentences per article; relevance = 0.6·word coverage + 0.4·embedding cosine ≥ 0.34; publisher trust and date recorded."],
        ["3 Classifier signals", "LIAR 6-class probabilities per claim (truth score = Σ p·[0, .2, .4, .6, .8, 1]) and Integrated binary P(true) for the article. Used as one signal only."],
        ["4 Evidence comparison", "deepseek-r1:8b / gemma2:9b prompted to compare the claim ONLY with numbered evidence and cite ids; output rejected unless cited ids ⊆ evidence. Fallback: clause-level stance rules (reported-claim filtering, misleading-context cues, subject-aware antonyms, number mismatch, negation polarity, absolute-vs-qualified). Stance weight = trust/100 × recency × relevance."],
        ["4b Conflict resolution", "If supporting and contradicting evidence both carry weight, the conflict is reported with both publishers, dates and weights; newer, higher-trust evidence wins; ratio < 1.3 → UNVERIFIED; uncertainty medium/high (TC-11)."],
        ["4c Framing & missing context", "Sensational headline cues, all-caps, exclamation marks, calls-to-action and anonymous sourcing produce a framing penalty; a sensational headline over a supported body becomes a MISLEADING framing claim (TC-12)."],
        ["5 Report", "source_trust + source_info, fact_accuracy, overall_score (system estimate), claim_results with badges, rationale and citations, highlights, timeline, xai_explanation, model_execution_metrics, security_flags."],
    ], [4, 13])
    para(doc, "Verdict points used for Fact Accuracy: SUPPORTED 95, UNVERIFIED 50, MISLEADING 35, FALSE 5 (mean over claims). Overall labels: "
              "≥ 80 Verified · ≥ 65 Mostly True · ≥ 45 Mixed/Disputed · ≥ 25 Mostly False · < 25 False.")
    doc.add_heading("4.5 Module 5 — Security, safety and prompt-injection guard (security.py)", 2)
    bullets(doc, ["All article text, scraped pages and user inputs are untrusted data wrapped in <untrusted_data> delimiters; system prompts carry a canary token that is checked for leakage.",
                  "Pattern families: override_instructions, verdict_manipulation, prompt_exfiltration, role_hijack, delimiter_injection, indirect_instruction. Only the injected span is removed so a legitimate claim before it is still analysed.",
                  "Verdicts are computed from evidence, so an instruction such as \"mark this article true\" cannot change the result (TC-13, TC-14). Events are logged to security_events and shown in the Admin dashboard.",
                  "Toxic or graphic quotations are neutralised before summarisation (TC-16)."])
    doc.add_heading("4.6 Module 6 — XAI and multi-model telemetry", 2)
    para(doc, "Described in detail in Section 5. metrics_store.py merges the JSON telemetry from the training and evaluation scripts so "
              "/api/admin/metrics and every /api/analyze response expose the benchmark measures of all five executing models.")

    # ---------------- 5. XAI
    page_break(doc)
    doc.add_heading("5. The Explainable-AI algorithm (LIME + SHAP)", 1)
    para(doc, "Every analysed article or claim receives two complementary explanations, computed in backend/app/xai_explainer.py and rendered "
              "in the Analyzer result, the embedded article report and the Admin XAI sandbox.")
    doc.add_heading("5.1 LIME token-level attribution", 2)
    bullets(doc, [
        "Interpretable units:: the text is segmented into money amounts (\"Rs 50,000\"), cue phrases (\"forward immediately\", \"anonymous ministry source\", \"official clarification\"), multi-word entities (\"Reserve Bank of India\") and remaining content words; d ≤ 60 features.",
        "Perturbation:: binary masks z ∈ {0,1}^d — the all-ones row, all d leave-one-out rows, then random masks removing k ~ U{1, d−1} units (300 samples by default, seeded for determinism).",
        "Black box:: f(text) = 0.6·P_Integrated(true) + 0.4·LIAR truth score — the supervised credibility signal.",
        "Locality kernel:: cosine distance to the original instance with an exponential kernel.",
        "Surrogate:: weighted ridge regression; the coefficients wᵢ ∈ [−1, +1] are the signed change in P(credible) contributed by each unit.",
    ])
    eq(doc, "D(z) = 1 − √(‖z‖₁ / d),     π(z) = exp(−D(z)² / σ²),  σ = 0.25")
    eq(doc, "ŵ = argmin_w Σ_z π(z)·( f(z) − w₀ − Σᵢ wᵢ zᵢ )² + α‖w‖²,  α = 1")
    para(doc, "Local fidelity is reported as the weighted R² of the surrogate. Negative weights (red pills) push towards False/Misleading; "
              "positive weights (green pills) towards Verified.")
    doc.add_heading("5.2 SHAP-style exact Shapley attribution over 5 pipeline dimensions", 2)
    para(doc, "The overall score is produced by the credibility aggregator v(x) over five dimensions — S Source Reliability, E RAG Evidence Match "
              "(= Fact Accuracy), L BERT-LIAR 6-class signal, B BERT ISOT-Bharat binary signal and F Framing/Omission penalty:")
    eq(doc, "v(x) = clip( 0.25·S + 0.40·E + 0.15·L + 0.20·B − F , 0, 100 ),   and v ≤ 38 whenever E < 30")
    para(doc, "Relative to the neutral prior x⁰ = (S, E, L, B, F) = (50, 50, 50, 50, 0) — baseline score 50 — the exact Shapley value of "
              "dimension i is computed over all 2⁵ = 32 coalitions:")
    eq(doc, "φᵢ = Σ_{T ⊆ N∖{i}}  |T|!(|N|−|T|−1)! / |N|!  ·  [ v(x_{T∪{i}}, x⁰_rest) − v(x_T, x⁰_rest) ]")
    para(doc, "Efficiency guarantees v(x⁰) + Σᵢ φᵢ = v(x): every point of the final score is attributed to exactly one dimension, including the "
              "non-linear contradiction cap, which plain weight×value attributions would mis-assign.")
    picture(doc, FIG / "shap_example.png", 14, "Figure 2 — SHAP contributions for the wireframe example (viral Rs 50,000 post, tb-001).")
    picture(doc, FIG / "lime_example.png", 14, "Figure 3 — Top LIME token weights for the same article.")
    doc.add_heading("5.3 Verification of the XAI implementation (scripts/verify_xai.py)", 2)
    ch = XAIV.get("checks", {})
    table(doc, ["Property", "Result"], [
        ["SHAP efficiency  |v(x⁰)+Σφ − v(x)| < 1e-6 on all demo articles", str(ch.get("shap_efficiency"))],
        ["SHAP null player  (dimension at its prior ⇒ φ = 0)", str(ch.get("shap_null_player"))],
        ["LIME mean local fidelity R²", f(ch.get("lime_mean_r2"))],
        ["Deletion test: removing top-3 negative tokens raises f", f(ch.get("lime_deletion_negative_raises"))],
        ["Deletion test: removing top-3 positive tokens lowers f", f(ch.get("lime_deletion_positive_lowers"))],
        ["Deterministic LIME weights across runs", str(ch.get("lime_deterministic"))],
    ], [12, 5])
    rows = [[r["article"], r["overall"], r["shap_baseline"], r["shap_sum"], r["lime_r2"], ", ".join(r["top_negative"])] for r in XAIV.get("rows", [])]
    table(doc, ["Article", "Overall", "φ₀", "Σφ", "LIME R²", "Top negative LIME tokens"], rows, [1.8, 1.6, 1.3, 1.6, 1.8, 8.9], font=8)
    doc.add_heading("5.4 How the XAI is displayed in the UI", 2)
    bullets(doc, ["LIME token heatmap:: the analysed text is rendered with colour-intensity pills (red = negative, green = positive) and the exact weight as a superscript; side lists rank the strongest tokens.",
                  "SHAP feature bars:: horizontal bars centred on zero for the five dimensions, preceded by the additive equation baseline ± φ = final and followed by the efficiency check.",
                  "Admin XAI sandbox:: free-text LIME with a perturbation-sample slider; five sliders that recompute local SHAP live; global SHAP (mean |φ| over all saved reports)."])

    # ---------------- 6. model performance
    page_break(doc)
    doc.add_heading("6. Deep-learning model performance measures", 1)
    para(doc, "All five executing models expose these measures via /api/admin/metrics; the compact benchmark badges also appear inside every "
              "/api/analyze response (model_execution_metrics) and are rendered in the Analyzer telemetry strip and the Admin dashboard.")
    r = PIPE.get("retriever", {})
    doc.add_heading("6.1 Model 1 — Sentence-BERT + FAISS + MMR retriever", 2)
    table(doc, ["Precision@5", "Recall@5", "MRR", "ILD", "Unique source ratio", "Latency (ms/query)", "Queries", "Back-end"],
          [[r.get("precision_at_k"), r.get("recall_at_k"), r.get("mrr"), r.get("ild"), r.get("unique_source_ratio"), r.get("latency_ms"),
            r.get("n_queries"), r.get("backend", "—")]], [1.8, 1.6, 1.4, 1.4, 2, 2, 1.6, 5.2], font=8)
    s = PIPE.get("summarizer", {})
    doc.add_heading("6.2 Model 2 — Qwen 4B summariser & intent extractor", 2)
    table(doc, ["ROUGE-1", "ROUGE-2", "ROUGE-L", "Faithfulness", "Entity & qty preservation", "Slot F1", "Articles", "Back-end"],
          [[s.get("rouge1"), s.get("rouge2"), s.get("rougeL"), s.get("faithfulness"), s.get("entity_quantity_preservation"),
            s.get("intent_slot_f1"), s.get("n"), s.get("backend", "—")]], [1.6, 1.6, 1.6, 1.9, 2.4, 1.5, 1.5, 4.9], font=8)
    m3 = CLF.get("liar", {}).get("primary", {})
    doc.add_heading("6.3 Model 3 — BERT-LIAR 6-class claim classifier", 2)
    table(doc, ["Accuracy", "Macro P", "Macro R", "Macro F1", "ROC-AUC (OvR)", "Test n", "Back-end"],
          [[m3.get("accuracy"), m3.get("macro_precision"), m3.get("macro_recall"), m3.get("macro_f1"), m3.get("roc_auc"), m3.get("n_test"), m3.get("backend")]], [2, 2, 2, 2, 2.4, 1.6, 5])
    table(doc, ["Class", "Precision", "Recall", "F1", "Support"],
          [[k, v["precision"], v["recall"], v["f1"], v["support"]] for k, v in m3.get("per_class", {}).items()], [4, 3, 3, 3, 3])
    picture(doc, FIG / "cm_liar.png", 9, "Figure 4 — 6×6 confusion matrix (test split). Errors concentrate between adjacent truthfulness grades.")
    m4 = CLF.get("integrated", {}).get("primary", {})
    doc.add_heading("6.4 Model 4 — BERT-Integrated (ISOT + BharatFakeNewsKosh) binary classifier", 2)
    table(doc, ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "Test n", "Back-end"],
          [[m4.get("accuracy"), m4.get("precision"), m4.get("recall"), m4.get("f1"), m4.get("roc_auc"), m4.get("n_test"), m4.get("backend")]], [2, 2, 2, 2, 2, 1.6, 5.4])
    picture(doc, FIG / "cm_integrated.png", 6.5, "Figure 5 — 2×2 confusion matrix (test split).")
    base_rows = []
    for n in ("liar", "integrated"):
        for b, mm in CLF.get(n, {}).get("baselines", {}).items():
            base_rows.append([n, b, mm.get("backend"), mm.get("accuracy"), mm.get("macro_f1"), mm.get("roc_auc")])
    table(doc, ["Dataset", "Model", "Back-end", "Accuracy", "Macro F1", "ROC-AUC"], base_rows, [2.5, 3, 4, 2.5, 2.5, 2.5])
    m5 = PIPE.get("reasoner", {})
    doc.add_heading("6.5 Model 5 — DeepSeek-R1 8B / Gemma2 9B evidence reasoner", 2)
    table(doc, ["Verdict acc.", "Hallucination", "Citation validity", "Context relevancy*", "Answer relevancy*", "Injection resistance", "5-run consistency", "Back-end"],
          [[m5.get("verdict_accuracy"), m5.get("hallucination_rate"), m5.get("citation_validity"), m5.get("contextual_relevancy"),
            m5.get("answer_relevancy"), m5.get("prompt_injection_resistance"), m5.get("five_run_consistency"), m5.get("backend", "—")]], [1.7, 1.8, 1.8, 2, 2, 2, 2, 3.7], font=8)
    para(doc, "* RAGAs-style relevancy values are embedding-based proxies in offline mode; with Ollama + deepeval/ragas installed the LLM-judged "
              "versions are used. The 24-claim benchmark covers 9 SUPPORTED, 5 FALSE, 3 MISLEADING and 7 UNVERIFIED claims.")

    # ---------------- 7. ablation
    page_break(doc)
    doc.add_heading("7. Section V — Ablation study (Table 1)", 1)
    para(doc, "scripts/run_ablation_study.py evaluates the complete system and six ablated variants on identical held-out samples "
              f"({json.dumps(ABL.get('samples', {}))}). Metrics of modules a configuration does not touch are shared with the complete system.")
    cols = [("precision_at_5", "P@5"), ("recall_at_5", "R@5"), ("mrr", "MRR"), ("ild", "ILD"), ("unique_source_ratio", "Uniq"),
            ("verdict_accuracy", "Verd."), ("hallucination_rate", "Hall."), ("citation_validity", "Cite"), ("credibility_band_accuracy", "Band"),
            ("misleading_recall", "Misl."), ("rouge1", "R-1"), ("rouge2", "R-2"), ("rougeL", "R-L"), ("factual_coverage", "Cov.")]
    table(doc, ["Configuration"] + [c for _, c in cols],
          [[r["configuration"].replace("Model ", "M").replace("w/o ", "−")] + [r[k] for k, _ in cols] for r in ABL.get("rows", [])],
          [3.6] + [0.95] * 14, font=6.5)
    picture(doc, FIG / "ablation.png", 16.5, "Figure 6 — Selected ablation metrics.")
    rows = ABL.get("rows", [])
    if len(rows) == 7:
        b = rows[0]
        bullets(doc, [
            f"Model 1 (−semantic retrieval):: P@5 {f(b['precision_at_5'])} → {f(rows[1]['precision_at_5'])}, MRR {f(b['mrr'])} → {f(rows[1]['mrr'])}; keyword matching misses paraphrased queries.",
            f"Model 2 (−personalisation):: P@5 {f(b['precision_at_5'])} → {f(rows[2]['precision_at_5'])}; profile-driven queries lose their topic prior.",
            f"Model 3 (−MMR):: ILD {f(b['ild'])} → {f(rows[3]['ild'])} and unique-source ratio {f(b['unique_source_ratio'])} → {f(rows[3]['unique_source_ratio'])} — the feed collapses onto fewer publishers.",
            f"Model 4 (−external evidence):: verdict accuracy {f(b['verdict_accuracy'])} → {f(rows[4]['verdict_accuracy'])}, hallucination {f(b['hallucination_rate'])} → {f(rows[4]['hallucination_rate'])}, citation validity → {f(rows[4]['citation_validity'])}.",
            f"Model 5 (−multi-dimensional credibility):: credibility band accuracy {f(b['credibility_band_accuracy'])} → {f(rows[5]['credibility_band_accuracy'])}; misleading-framing recall {f(b['misleading_recall'])} → {f(rows[5]['misleading_recall'])}.",
            f"Model 6 (−LLM summarisation):: ROUGE-L {f(b['rougeL'])} → {f(rows[6]['rougeL'])}, factual coverage {f(b['factual_coverage'])} → {f(rows[6]['factual_coverage'])} with lead-3 extraction.",
        ])
    para(doc, "Back-ends: " + ", ".join(f"{k} = {v}" for k, v in ABL.get("backends", {}).items()))

    # ---------------- 8. tests
    page_break(doc)
    doc.add_heading("8. Section VI — 18 LLM test cases (Tables 2–5)", 1)
    para(doc, f"scripts/run_llm_test_suite.py (also tests/test_deepeval_suite.py and the Admin ▶ button) executed {LLM.get('total', 18)} cases; "
              f"{LLM.get('passed', '—')} passed. Engine: {json.dumps(LLM.get('engine', {}))}.")
    titles = {"table2_L1_component": "Table 2 — L1 component tests", "table3_L2_rag": "Table 3 — L2 RAG / evidence-review tests",
              "table4_L3_L4_safety_e2e": "Table 4 — L3 safety & robustness and L4 end-to-end tests"}
    for key, title in titles.items():
        doc.add_heading(title, 2)
        table(doc, ["TC", "Scenario", "Metric", "Target", "Measured", "Status"],
              [[c["id"], c["name"], c["metric"], c["target"], str(c["measured"]), c["status"]] for c in LLM.get("tables", {}).get(key, [])],
              [1.4, 5.2, 4, 2.4, 2.8, 1.4], font=7.5)
    doc.add_heading("Table 5 — Result-reporting template", 2)
    table(doc, ["Metric", "Target", "Measured", "Status"], [[r["metric"], r["target"], str(r["measured"]), r["status"]] for r in LLM.get("table5", [])],
          [8.5, 2.5, 3, 2.5])
    para(doc, f"Citation traceability (TC-18): {LLM.get('citation_traceability', '—')} of cited URLs belong to the retrieved evidence set, and every summary links to its original article.")
    doc.add_heading("8.1 What each test case checks", 2)
    for c in LLM.get("cases", []):
        para(doc, f" {c['name']}. Metric: {c['metric']} (target {c['target']}); measured {c['measured']} → {c['status']}.", c["id"])

    # ---------------- 9. UI guide
    page_break(doc)
    doc.add_heading("9. User-interface guide", 1)
    doc.add_heading("9.1 Header (all screens)", 2)
    bullets(doc, ["Utility strip:: date · E-Paper · EN / NE · ☀️/🌙 Dark/Light toggle (persisted) · Active Models pill (green dot = Ollama live, amber = offline fallback).",
                  "Brand bar:: TB logo, semantic search box (Enter runs Sentence-BERT/FAISS + MMR search with your profile), ↻ Live RSS Scrape, ✓ Check News.",
                  "Category bar:: News | Politics | Society | Economy | World | Sports | Opinion | Entertainment | Tech | ⚡ Fact-Check | 💬 AI Chatbot | 📊 Admin: Model Metrics & XAI | 👤 Profile.",
                  "BREAKING ticker:: scrolling headlines (Parliament session postponed — EC publishes voter roll update — Monsoon alert in Koshi/Kerala …)."])
    doc.add_heading("9.2 Screen 1 — Fake-News Analyzer (⚡ Fact-Check)", 2)
    bullets(doc, ["Stepper 1 Paste input → 2 Source checked → 3 Claims split → 4 Tagged report tracks progress.",
                  "Left panel:: Link | Text | Image tabs, Reasoning LLM selector (deepseek-r1:8b, gemma2:9b, llama3.1:8b; '(fallback engine)' when not pulled), input, ⚡ Analyze and eight 1-click demo presets (viral Rs 50,000 post, RBI SHOCKER headline, ISRO alien rumour, Chitradurga conflict, PM-KISAN omitted context, reused flood photo, prompt-injection attempt, no-evidence claim).",
                  "Right panel:: Source Trust and Fact Accuracy meters, overall estimate badge, injection notice, WHAT IS WRONG list with FALSE (red) / MISLEADING (amber) / UNVERIFIED (grey) / SUPPORTED (green) badges, rationale and evidence links, and Open report.",
                  "Below:: the 🧠 XAI box (LIME heatmap + SHAP bars) and the 📈 executing-models telemetry strip (back-end, latency, confidence, benchmark badges for all models).",
                  "Recent public checks:: three-column grid with colour-coded top bars."])
    doc.add_heading("9.3 Screen 2 — Article landing page with embedded report", 2)
    bullets(doc, ["Breadcrumb (Home / Economy / Rs 50,000 allowance claim), category pill, status pill (! Disputed N or ✓ Verified N), headline and byline with the original URL.",
                  "Pipeline A summary card:: summariser back-end, ROUGE-L and faithfulness badges, entity preservation, 'no verdict ✓', key points, attribution and a 1–5★ faithfulness rating stored in the feedback table.",
                  "Article body:: claims are highlighted in their verdict colour with the rationale on hover; the top factual claim is shown as a pull-quote.",
                  "Embedded Pipeline B report:: overall badge (e.g. 'False — 14/100'), meters with source summary ('Unknown outlet • no masthead • … • 12 days old') and claim counts, tagged claims, non-checked sentences, XAI panel, telemetry, Sources cited cards, Key highlights, Timeline of claim, Related stories (Verified / Disputed / Unreviewed).",
                  "Sticky sidebar:: IN THIS ARTICLE claim anchors, SOURCE BOX (outlet, trust, owner, masthead, corrections record, history, contact) and CHECK ANOTHER CLAIM."])
    doc.add_heading("9.4 Screen 3 — Admin: Model Metrics & XAI", 2)
    bullets(doc, ["Five model cards with every metric of Section 6, per-class F1 bars, 6×6, 2×2 and verdict confusion matrices, and baselines.",
                  "Interactive Table 1:: cells worse than the complete system are red, better are green, with deltas; click a column header to highlight it.",
                  "Table 5 and all 18 test cases with PASS/FAIL badges; click a row to see the raw details JSON.",
                  "▶️ Re-Run All Model Tests & Ablation Suite (or only classifiers / ablation / llm) starts a background job; the log streams in the job card and all tables refresh on completion.",
                  "Interactive XAI sandbox, security events, feedback ratings and evaluation history."])
    doc.add_heading("9.5 Screen 4 — AI Chatbot", 2)
    bullets(doc, ["Each answer shows extracted intent slots (intent, topic, entities, source, time, model) and clickable recommended articles with relevance and personalisation reasons plus a short Pipeline A summary.",
                  "Live ablation toggles:: Semantic retrieval, Personalisation and MMR checkboxes change the next query's ranking.",
                  "Fact-check questions are routed to Pipeline B and linked; medical-dosage and other out-of-scope requests are declined."])
    doc.add_heading("9.6 Profile", 2)
    para(doc, "Choose the active user (demo or guest/sparse), preferred topics and sources, followed entities, and review recent conversational queries.")

    # ---------------- 10. API
    page_break(doc)
    doc.add_heading("10. REST API reference", 1)
    table(doc, ["Method & path", "Purpose"], [
        ["GET /api/health", "status, article count, Ollama status, back-ends"],
        ["GET /api/models", "configured models, availability, fallback note"],
        ["GET /api/articles?topic=&q=", "feed (with latest report status) or semantic search"],
        ["GET /api/articles/{id}", "article + Pipeline A summary + latest report + related + source"],
        ["POST /api/search", "query, topic, source, time_window, use_semantic, use_personalization, use_mmr, mmr_lambda"],
        ["POST /api/chat", "conversational intent → reply, slots, articles (or decline / clarification / fact-check link)"],
        ["POST /api/summarize", "Pipeline A summary of article_id or pasted text"],
        ["POST /api/analyze", "Pipeline B report: input_type link | text | image | article, reasoning_model, instruction, lime_samples"],
        ["GET /api/reports/recent, /api/reports/{id}", "saved credibility reports"],
        ["POST /api/scrape/rss, /api/scrape/url", "live ingestion"],
        ["GET|PUT /api/profile/{user_id}", "user preferences"],
        ["POST /api/feedback", "1–5 rating for summary / report / recommendation / chat"],
        ["POST /api/xai/explain, GET /api/xai/global", "XAI sandbox and global SHAP"],
        ["GET /api/admin/metrics", "all 5 models + ablation + tests + history + security + feedback"],
        ["POST /api/admin/run-tests, GET /api/admin/jobs/{id}", "re-run classifiers / ablation / llm suites in the background"],
        ["GET /api/security/events", "prompt-injection and safety log"],
    ], [6, 11])
    para(doc, "Interactive OpenAPI documentation is served at http://localhost:8000/docs.")

    # ---------------- 11. commands
    page_break(doc)
    doc.add_heading("11. Installation and exact terminal commands", 1)
    doc.add_heading("11.1 Windows (PowerShell)", 2)
    code(doc, r"""
cd TrendyBits
py -3.11 -m venv .venv
.\.venv\Scripts\Activate.ps1
pip install -r backend\requirements.txt
# optional deep-learning stack (BERT fine-tuning, SBERT, FAISS, DeepEval, RAGAs):
pip install --extra-index-url https://download.pytorch.org/whl/cpu -r backend\requirements-ml.txt
python scripts\run_all.py                          # datasets → classifiers → XAI check → ablation → 18 tests
cd backend; uvicorn app.main:app --reload --port 8000
# second terminal
cd TrendyBits\frontend; npm install; npm run dev     # http://localhost:5173
""")
    doc.add_heading("11.2 macOS / Linux (bash / zsh)", 2)
    code(doc, """
cd TrendyBits
python3.11 -m venv .venv && source .venv/bin/activate
pip install -r backend/requirements.txt
pip install --extra-index-url https://download.pytorch.org/whl/cpu -r backend/requirements-ml.txt   # optional
python scripts/run_all.py
(cd backend && uvicorn app.main:app --reload --port 8000) &
cd frontend && npm install && npm run dev
""")
    doc.add_heading("11.3 Ollama dual-LLM setup", 2)
    code(doc, """
# install from https://ollama.com/download, then:
ollama pull qwen:4b            # Pipeline A (or: ollama pull qwen2.5:3b)
ollama pull deepseek-r1:8b     # Pipeline B (or: ollama pull gemma2:9b / llama3.1:8b)
ollama pull mistral:7b         # LLM-as-a-judge for DeepEval
ollama serve                   # if not already running as a service
python scripts/check_ollama.py # verifies each model with a smoke prompt
# switch models without code changes:
export TRENDYBITS_SUMMARIZER_MODEL=qwen2.5:3b TRENDYBITS_REASONER_MODEL=gemma2:9b     # bash
$env:TRENDYBITS_REASONER_MODEL="gemma2:9b"                                            # PowerShell
export TRENDYBITS_USE_OLLAMA=0   # force the deterministic fallback engine
""")
    doc.add_heading("11.4 IEEE Dataport ingestion", 2)
    code(doc, """
python scripts/import_ieee_dataport.py --newssumm path/to/NewsSumm.xlsx
python scripts/import_ieee_dataport.py --liar train.tsv valid.tsv test.tsv
python scripts/import_ieee_dataport.py --isot-true True.csv --isot-fake Fake.csv --bharat bharat_fakenewskosh_en.csv
python scripts/import_ieee_dataport.py --newssumm NewsSumm.csv --load-into-db 200   # also add articles to the feed
python scripts/import_ieee_dataport.py --generate                                    # synthetic partitions
""")
    doc.add_heading("11.5 Training, evaluation and tests", 2)
    code(doc, """
python scripts/train_bert_classifier.py --mode fast                     # TF-IDF + LR and MLP baselines
python scripts/train_bert_classifier.py --mode bert --epochs 3 --batch-size 16 --max-len 128
python scripts/train_bert_classifier.py --mode bert --dataset liar --sample 2000   # quicker CPU run
python scripts/verify_xai.py
python scripts/run_ablation_study.py --summaries 70
python scripts/run_llm_test_suite.py
pytest -q tests
pytest tests/test_deepeval_suite.py -v
deepeval test run tests/test_deepeval_suite.py      # if deepeval is installed
python scripts/build_manual.py                       # regenerate this manual
""")
    doc.add_heading("11.6 Docker", 2)
    code(doc, """
docker compose up --build                          # API :8000, UI :5173, uses host Ollama if running
docker compose --profile ollama up --build         # + Ollama container that pulls qwen:4b, deepseek-r1:8b, mistral:7b
INSTALL_ML=true docker compose build backend       # bake torch/transformers/SBERT/FAISS into the image
docker compose exec backend python /app/scripts/run_llm_test_suite.py
docker compose down -v
""")

    # ---------------- 12. troubleshooting
    doc.add_heading("12. Troubleshooting and known limitations", 1)
    table(doc, ["Symptom", "Fix"], [
        ["'Active Models … offline fallback' in the header", "Start Ollama and pull the models (11.3); the UI switches automatically within 30 s."],
        ["UI shows 'is the backend running on :8000?'", "Start uvicorn from backend/; if the API runs elsewhere set VITE_API_PROXY before npm run dev."],
        ["Live RSS scrape reports feed errors", "No internet or the publisher changed its feed URL; edit RSS_FEEDS in backend/app/scraper.py."],
        ["--mode bert exits with an import error", "pip install -r backend/requirements-ml.txt (CPU wheels via the PyTorch index)."],
        ["Windows console shows garbled symbols", "All scripts force UTF-8 output; use Windows Terminal or set PYTHONIOENCODING=utf-8."],
    ], [6, 11])
    bullets(doc, ["The deterministic NLI fallback is rule-based: it handles negation, antonyms with subject awareness, numeric mismatches, reported speech, reused-media cues and omitted qualifiers, but not deep paraphrase or multi-hop reasoning — use an Ollama reasoner for that.",
                  "Image input checks the caption/context against evidence; no pixel-level image forensics is performed.",
                  "Credibility scores are system estimates for triage and must be read together with the cited evidence.",
                  "Shipped metrics come from synthetic partitions and a small demo corpus (Section 1.1); re-run on the IEEE Dataport data before reporting results."])

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build()
    print(f"manual written → {out} ({out.stat().st_size / 1024:.0f} KB)")
